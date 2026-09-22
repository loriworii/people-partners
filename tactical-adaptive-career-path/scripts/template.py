"""
Tactical + Adaptive Career Path — Python Template
==================================================
Copy this file and fill in the CONTENT section.
Do not change the BOILERPLATE sections — they encode the validated visual spec.

Usage:
  1. Set OUT to your output path
  2. Define your LEVELS list
  3. Fill in RESP (responsibilities) and SKILLS dictionaries
  4. Set BEHAVIOURS from source file or standard set
  5. Run: python3 this_file.py
  6. Validate: python3 /path/to/recalc.py "output.xlsx"
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# ── OUTPUT PATH — change this ───────────────────────────────────────────────
OUT = "/sessions/REPLACE_SESSION_ID/mnt/outputs/[Team Name] Career Path - Tactical + Adaptive.xlsx"

# ══════════════════════════════════════════════════════════════════════════════
# BOILERPLATE — do not change
# ══════════════════════════════════════════════════════════════════════════════

FILL_HEADER  = PatternFill("solid", fgColor="C1FF14")   # lime green
FILL_SECTION = PatternFill("solid", fgColor="FF62E6")   # pink
FILL_TITLE   = PatternFill("solid", fgColor="EFEFEF")   # grey
FILL_WHITE   = PatternFill("solid", fgColor="FFFFFF")

def fnt(bold=False, size=9, name="DM Sans"):
    return Font(bold=bold, size=size, name=name)

def aln(h="left", v="center", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def style(cell, fill, font, alignment):
    cell.fill = fill; cell.font = font; cell.alignment = alignment

def fmt(ownership, tactical, adaptive, impact):
    """Build a T+A responsibility cell."""
    t  = f"Ownership: {ownership}\n\n"
    t += "Tactical (Execute):\n"
    for b in tactical: t += f"* {b}\n"
    t += "\nAdaptive (Evolve):\n"
    for b in adaptive: t += f"* {b}\n"
    t += f"\nImpact: {impact}"
    return t.strip()

def skl(bullets, signal):
    """Build a skills cell with Observable Signal."""
    return "\n".join(f"* {b}" for b in bullets) + f"\n\nObservable signal: {signal}"

# ══════════════════════════════════════════════════════════════════════════════
# CONTENT — fill this in
# ══════════════════════════════════════════════════════════════════════════════

# Levels: list of (code, title) tuples — these map to content columns C, D, E...
# Examples:
#   IC track: [("IC2","Associate Data Analyst"), ("IC3","Data Analyst"), ...]
#   M track:  [("M2","Manager, External Comms"), ("M3","Senior Manager, External Comms"), ...]
LEVELS = [
    ("IC2", "Associate [Role]"),
    ("IC3", "[Role]"),
    ("IC4", "Senior [Role]"),
    ("IC5", "Staff [Role]"),
    ("IC6", "Sr Staff [Role]"),
    # ("IC7", "Principal [Role]"),   # add/remove levels as needed
]

# Column keys matching LEVELS order
KEYS = [get_column_letter(3 + i) for i in range(len(LEVELS))]
COLS = [3 + i for i in range(len(LEVELS))]
LAST_COL = get_column_letter(2 + len(LEVELS))

# Responsibility domain labels (5 required)
RESP_LABELS = [
    "Domain 1 — replace me",
    "Domain 2 — replace me",
    "Domain 3 — replace me",
    "Domain 4 — replace me",
    "Domain 5 — replace me",
]

# Responsibility content: dict keyed by row (3–7), each value is dict keyed by column letter
# Use fmt() for each cell.
# RESP[3] = row 3 = first responsibility domain
# RESP[3]["C"] = content for IC2 (or M2), first responsibility domain
RESP = {
    3: {key: fmt(
        f"Ownership statement for level {key} in domain 1.",
        ["Tactical bullet 1", "Tactical bullet 2", "Tactical bullet 3"],
        ["Adaptive bullet 1", "Adaptive bullet 2"],
        "Impact statement."
    ) for key in KEYS},
    4: {key: fmt(
        f"Ownership statement for level {key} in domain 2.",
        ["Tactical bullet 1", "Tactical bullet 2"],
        ["Adaptive bullet 1", "Adaptive bullet 2"],
        "Impact statement."
    ) for key in KEYS},
    5: {key: fmt(
        f"Ownership statement for level {key} in domain 3.",
        ["Tactical bullet 1", "Tactical bullet 2"],
        ["Adaptive bullet 1", "Adaptive bullet 2"],
        "Impact statement."
    ) for key in KEYS},
    6: {key: fmt(
        f"Ownership statement for level {key} in domain 4.",
        ["Tactical bullet 1", "Tactical bullet 2"],
        ["Adaptive bullet 1", "Adaptive bullet 2"],
        "Impact statement."
    ) for key in KEYS},
    7: {key: fmt(
        f"Ownership statement for level {key} in domain 5.",
        ["Tactical bullet 1", "Tactical bullet 2"],
        ["Adaptive bullet 1", "Adaptive bullet 2"],
        "Impact statement."
    ) for key in KEYS},
}

# Behaviours: list of (name, lower_tier_text, upper_tier_text)
# Lower tier = IC2–IC4 or equivalent; Upper tier = IC5+ or all managers
# Copy from source file verbatim. Use standard set from references/framework.md if no source.
BEHAVIOUR_SPLIT_INDEX = 3  # first N columns (0-indexed) use lower tier; rest use upper tier
BEHAVIOURS = [
    (
        "We act with kindness",
        "* Treat others with care, compassion, and dignity in daily interactions.\n"
        "* Regularly check in on team members and provide thoughtful, empathetic feedback.\n"
        "* Foster inclusivity and belonging by listening deeply and understanding diverse perspectives.\n"
        "* Collaborate with transparency, consulting with impacted stakeholders before making decisions.",
        "* Proactively lead initiatives that promote kindness and empathy within the team.\n"
        "* Mentor junior team members, modelling a culture of compassion and authentic communication.\n"
        "* Ensure transparency and inclusiveness in decision-making, involving relevant stakeholders.\n"
        "* Advocate for a supportive and caring community, recognising achievements and contributions.",
    ),
    (
        "We cultivate community",
        "* Support and empower team members by leveraging each other's strengths to achieve goals.\n"
        "* Actively participate in team efforts, contributing to open, inclusive decision-making.\n"
        "* Recognise and value diversity in the workplace, fostering collaboration across different backgrounds.\n"
        "* Build partnerships with internal and external stakeholders to align with shared goals.",
        "* Take ownership of cultivating a strong team culture, emphasising collective success over individual achievement.\n"
        "* Lead cross-functional collaborations that harness diverse expertise for innovative solutions.\n"
        "* Drive inclusive decision-making processes that respect and integrate diverse perspectives.\n"
        "* Build strategic partnerships with key internal and external stakeholders to further team and company goals.",
    ),
    (
        "We continually evolve",
        "* Approach all projects with a growth mindset, learning from mistakes and seeking continuous improvement.\n"
        "* Provide and accept feedback openly, encouraging personal and team development.\n"
        "* Embrace change and actively seek new learning opportunities to refine your skills and knowledge.\n"
        "* Work with agility, adapting quickly to new insights or data from the market or team.",
        "* Lead by example, modelling continuous learning and development within the team.\n"
        "* Proactively solicit and integrate diverse feedback to drive team evolution.\n"
        "* Guide the team through changes, leveraging market insights to continuously refine strategies.\n"
        "* Mentor others in adopting a growth mindset, fostering a culture of iteration and learning.",
    ),
    (
        "We embrace innovation",
        "* Encourage creative problem-solving by seeking diverse opinions and exploring new solutions.\n"
        "* Actively contribute to innovative discussions, asking questions and sharing knowledge with the team.\n"
        "* Adapt to changes and emerging technologies, staying agile in your work approach.\n"
        "* Participate in experimentation and continuous improvement to drive team innovation.",
        "* Drive innovative projects, fostering a culture of curiosity and creative thinking.\n"
        "* Proactively lead teams through experimentation and agile development processes.\n"
        "* Advocate for new technologies and processes that will enable the team to stay ahead of industry trends.\n"
        "* Inspire others to embrace change and position the team at the forefront of innovation.",
    ),
    (
        "We deliver impact",
        "* Set clear objectives for your work, ensuring alignment with team and company goals.\n"
        "* Follow through on commitments and be accountable for delivering high-quality results.\n"
        "* Identify opportunities to deliver impactful outcomes through collaboration and thoughtful planning.\n"
        "* Participate in swift, collective decision-making that drives progress.",
        "* Lead strategic initiatives that deliver significant impact, driving team and company success.\n"
        "* Ensure the team sets measurable goals and is accountable for delivering exceptional outcomes.\n"
        "* Align resources and efforts toward delivering high-quality strategies and solutions.\n"
        "* Evaluate results rigorously, iterating processes and strategies to improve impact continuously.",
    ),
]

# Skills: list of (domain_label, [content_per_level])
# Each level content is built with skl(bullets_list, signal_string)
# 6 domains required; content list must have one entry per level
SKILL_LEVEL_LABELS = ["Foundational", "Intermediate", "Advanced", "Expert", "Strategic", "Visionary"]
# For manager tracks use: ["Manager", "Senior Manager", "Associate Director", "Director", "Senior Director", "VP"]

SKILLS = [
    (
        "Skill Domain 1 — replace",
        [skl(["Bullet 1", "Bullet 2"], "Observable signal for this level") for _ in LEVELS],
    ),
    (
        "Skill Domain 2 — replace",
        [skl(["Bullet 1", "Bullet 2"], "Observable signal for this level") for _ in LEVELS],
    ),
    (
        "Skill Domain 3 — replace",
        [skl(["Bullet 1", "Bullet 2"], "Observable signal for this level") for _ in LEVELS],
    ),
    (
        "Skill Domain 4 — replace",
        [skl(["Bullet 1", "Bullet 2"], "Observable signal for this level") for _ in LEVELS],
    ),
    (
        "Skill Domain 5 — replace",
        [skl(["Bullet 1", "Bullet 2"], "Observable signal for this level") for _ in LEVELS],
    ),
    (
        "Skill Domain 6 — replace",
        [skl(["Bullet 1", "Bullet 2"], "Observable signal for this level") for _ in LEVELS],
    ),
]

# Career path description — shown in the title row (Row 2)
# Keep this concise — it should name the function/team scope and explain Tactical vs Adaptive
# in one or two sentences. See the reference for length guidance.
CAREER_PATH_DESCRIPTION = (
    "Career Path — Tactical + Adaptive Performance\n"
    "[Replace with function scope, e.g. 'Covers PR, Corporate Comms and Product Comms.'] "
    "Each level defines what must be executed reliably (Tactical) and where the individual is "
    "expected to go beyond the defined scope (Adaptive). Useful for performance management, "
    "promotion decisions and individual development."
)

# ══════════════════════════════════════════════════════════════════════════════
# BUILD WORKBOOK — do not change below this line
# ══════════════════════════════════════════════════════════════════════════════

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "[Team] Career Path"  # update sheet name

# Row 1: level headers
ws.row_dimensions[1].height = 60.0
style(ws["A1"], FILL_WHITE, fnt(size=9), aln())
ws["B1"].value = f"{LEVELS[0][0]} > {LEVELS[-1][0]}"
style(ws["B1"], FILL_WHITE, fnt(size=8), aln(h="center"))
for i, (code, title) in enumerate(LEVELS):
    cell = ws.cell(row=1, column=COLS[i])
    cell.value = f"{code}\n{title}"
    style(cell, FILL_HEADER, fnt(bold=True, size=10), aln(h="center"))

# Row 2: career path description (merged)
ws.row_dimensions[2].height = 37.5
ws.merge_cells("A2:B2")
ws["A2"] = CAREER_PATH_DESCRIPTION
style(ws["A2"], FILL_TITLE, fnt(bold=True, size=9), aln(wrap=True))
merge_range = f"C2:{LAST_COL}2"
ws.merge_cells(merge_range)
ws["C2"] = CAREER_PATH_DESCRIPTION
style(ws["C2"], FILL_TITLE, fnt(bold=True, size=9), aln(h="center", wrap=True))

# Rows 3–7: responsibilities
ws.merge_cells("A3:A7")
ws["A3"] = "Responsibilities"
style(ws["A3"], FILL_SECTION, fnt(bold=True, size=10), aln(h="center"))
for row, label in zip(range(3, 8), RESP_LABELS):
    ws.row_dimensions[row].height = 240
    b = ws.cell(row=row, column=2); b.value = label
    style(b, FILL_WHITE, fnt(bold=True, size=9), aln())
    for col, key in zip(COLS, KEYS):
        cell = ws.cell(row=row, column=col)
        cell.value = RESP[row][key]
        style(cell, FILL_WHITE, fnt(size=9), aln(wrap=True))

# Rows 8–12: behaviours
ws.merge_cells("A8:A12")
ws["A8"] = "Behaviours"
style(ws["A8"], FILL_SECTION, fnt(bold=True, size=10), aln(h="center"))
for r_idx, (beh_name, lower, upper) in enumerate(BEHAVIOURS):
    row = 8 + r_idx
    ws.row_dimensions[row].height = 114.75
    b = ws.cell(row=row, column=2); b.value = beh_name
    style(b, FILL_WHITE, fnt(bold=True, size=9), aln())
    for c_idx, col in enumerate(COLS):
        cell = ws.cell(row=row, column=col)
        cell.value = lower if c_idx < BEHAVIOUR_SPLIT_INDEX else upper
        style(cell, FILL_WHITE, fnt(size=9), aln(wrap=True))

# Row 13: skills header
ws.row_dimensions[13].height = 37.5
ws.merge_cells("A13:B13")
ws["A13"] = (
    "Skills & Knowledge Matrix\n"
    "What good looks like at each level — with observable signals to make assessment "
    "transparent and consistent."
)
style(ws["A13"], FILL_TITLE, fnt(bold=True, size=9), aln(wrap=True))
for i, (col, label) in enumerate(zip(COLS, SKILL_LEVEL_LABELS[:len(LEVELS)])):
    cell = ws.cell(row=13, column=col)
    cell.value = label
    style(cell, FILL_TITLE, fnt(bold=True, size=9), aln(h="center"))

# Rows 14–19: skills
ws.merge_cells("A14:A19")
ws["A14"] = "Skills & Knowledge Matrix"
style(ws["A14"], FILL_SECTION, fnt(bold=True, size=10), aln(h="center"))
for s_idx, (domain, level_content) in enumerate(SKILLS):
    row = 14 + s_idx
    ws.row_dimensions[row].height = 165
    b = ws.cell(row=row, column=2); b.value = domain
    style(b, FILL_WHITE, fnt(bold=True, size=9), aln())
    for c_idx, col in enumerate(COLS):
        cell = ws.cell(row=row, column=col)
        cell.value = level_content[c_idx]
        style(cell, FILL_WHITE, fnt(size=9), aln(wrap=True))

# Column widths
ws.column_dimensions["A"].width = 33.0
ws.column_dimensions["B"].width = 45.0
for letter in KEYS:
    ws.column_dimensions[letter].width = 63.0

wb.save(OUT)
print(f"Saved: {OUT}")
